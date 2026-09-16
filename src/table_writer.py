"""表格输出：CSV / Markdown / XLSX / JSON。

表格口径
--------
* 默认布局：**行 = 元素**，**列 = 月份**（可 ``table.transpose`` 转置）。
* 温度的"年"值取有效月均值，降水/降水日数的"年"值取有效月累加。
* 单位随元素类型自动附加（温度 °C/°F，降水 毫米/英寸）。
* 缺失值统一使用 ``table.missing_placeholder``。
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Optional

from .models import CityClimate, format_number, normalize_series_key

RAIN_KEYS = ("rainfall", "raindays")
TEMP_KEYS = ("minTemp", "maxTemp", "meanTemp")


# ---- 文案与单位 --------------------------------------------------------

def series_label(city: CityClimate, key: str, cfg: dict[str, Any]) -> str:
    """元素显示名：rainfall 走 raintype 感知命名。"""
    key = normalize_series_key(key)
    table_cfg = cfg["table"]
    label = (table_cfg.get("row_labels") or {}).get(key) or ""
    series_cfg = (cfg.get("series") or {}).get(key) or {}
    if not label or label == "auto":
        if key == "rainfall":
            label = city.rain_label(sentence=True)
        else:
            label = series_cfg.get("label") if series_cfg.get("label") not in (None, "", "auto") else key
    if not table_cfg.get("unit_in_label", True):
        return str(label)
    return f"{label} ({unit_label(city, key, cfg)})"


def unit_label(city: CityClimate, key: str, cfg: dict[str, Any]) -> str:
    key = normalize_series_key(key)
    if key in TEMP_KEYS:
        return "°F" if (cfg["data"].get("temp_unit") or "C").upper() == "F" else "°C"
    if key == "rainfall":
        return "英寸" if (cfg["data"].get("rain_unit") or "mm").lower() == "inch" else city.rain_unit_label()
    if key == "raindays":
        return "天"
    return ""


def _precision(cfg: dict[str, Any], key: str) -> int:
    prec = cfg["table"].get("precision") or {}
    if key == "annual":
        return int(prec.get("annual", prec.get("temp", 1)))
    key = normalize_series_key(key)
    if key in ("rainfall", "raindays"):
        return int(prec.get(key, 1))
    return int(prec.get("temp", 1))


def _month_headers(city: CityClimate, cfg: dict[str, Any]) -> list[str]:
    style = cfg["data"].get("month_label_style", "1月")
    headers = city.month_labels(style)
    if cfg["table"].get("include_annual", cfg["data"].get("include_annual", False)):
        headers = headers + [cfg["table"].get("annual_label", "年")]
    return headers


def _row_values(city: CityClimate, key: str, cfg: dict[str, Any]) -> list[Optional[float]]:
    temp_unit = (cfg["data"].get("temp_unit") or "C").upper()
    rain_unit = (cfg["data"].get("rain_unit") or "mm").lower()
    values = city.values(key, temp_unit, rain_unit)
    if cfg["table"].get("include_annual", cfg["data"].get("include_annual", False)):
        values = values + [city.annual(key, temp_unit, rain_unit)]
    return values


def build_matrix(city: CityClimate, cfg: dict[str, Any]) -> tuple[list[str], list[str], list[list[str]]]:
    """返回 (表头, 行标签, 每个单元格的字符串值)。"""
    table_cfg = cfg["table"]
    missing = table_cfg.get("missing_placeholder", "—")
    rows = [normalize_series_key(k) for k in table_cfg.get("rows") or []]
    month_headers = _month_headers(city, cfg)

    if table_cfg.get("transpose"):
        # 转置：月份成为行，元素成为列（年份作为最后一行）
        labels = month_headers
        header = ["月份"] + [series_label(city, k, cfg) for k in rows]
        temp_unit = (cfg["data"].get("temp_unit") or "C").upper()
        rain_unit = (cfg["data"].get("rain_unit") or "mm").lower()
        use_annual = bool(table_cfg.get("include_annual", cfg["data"].get("include_annual", False)))
        cells = []
        for idx in range(len(month_headers)):
            line: list[str] = []
            for key in rows:
                if idx < len(city.months):
                    value = city.months[idx].value(key, temp_unit, rain_unit)
                    line.append(format_number(value, _precision(cfg, key), missing))
                elif use_annual:
                    line.append(format_number(city.annual(key, temp_unit, rain_unit),
                                              _precision(cfg, "annual"), missing))
                else:
                    line.append(missing)
            cells.append(line)
        return header, labels, cells

    header = ["月份"] + month_headers
    labels: list[str] = []
    cells = []
    for key in rows:
        labels.append(series_label(city, key, cfg))
        prec = _precision(cfg, key)
        cells.append([format_number(v, prec, missing) for v in _row_values(city, key, cfg)])
    return header, labels, cells


def build_notes(city: CityClimate, cfg: dict[str, Any]) -> list[str]:
    """脚注：统计时段与数据来源。"""
    table_cfg = cfg["table"]
    notes: list[str] = []
    if table_cfg.get("show_period_note", True):
        period = city.period_note()
        if period:
            note = (table_cfg.get("note_template") or "").format(
                period=period,
                temp_unit="°F" if (cfg["data"].get("temp_unit") or "C").upper() == "F" else "°C",
                rain_unit=city.rain_unit_label(),
            )
            notes.append(note.strip())
    if table_cfg.get("show_source_note", True):
        notes.append(city.source_note())
    return [n for n in notes if n]


def table_title(city: CityClimate, cfg: dict[str, Any]) -> str:
    return (cfg["table"].get("title_template") or "{city} 气候统计").format(
        city=city.city_name, city_id=city.city_id, member=city.member.mem_name
    )


# ---- 各格式写出 --------------------------------------------------------

def write_csv(path: Path, city: CityClimate, cfg: dict[str, Any]) -> Path:
    encoding = cfg["output"].get("csv_encoding", "utf-8-sig")
    header, labels, cells = build_matrix(city, cfg)
    rows = list(zip(labels, cells))
    with path.open("w", encoding=encoding, newline="") as fh:
        writer = csv.writer(fh)
        if cfg["table"].get("show_title_row", True):
            writer.writerow([table_title(city, cfg)])
            for note in build_notes(city, cfg):
                writer.writerow([note])
            writer.writerow([])
        writer.writerow(header)
        for label, values in rows:
            writer.writerow([label, *values])
    return path


def write_markdown(path: Path, city: CityClimate, cfg: dict[str, Any]) -> Path:
    header, labels, cells = build_matrix(city, cfg)
    lines: list[str] = []
    if cfg["table"].get("show_title_row", True):
        lines.append(f"### {table_title(city, cfg)}")
        for note in build_notes(city, cfg):
            lines.append(f"> {note}")
        lines.append("")
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for label, values in zip(labels, cells):
        lines.append("| " + " | ".join([label, *values]) + " |")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_json(path: Path, city: CityClimate, cfg: dict[str, Any]) -> Path:
    payload = city.to_dict()
    payload["notes"] = build_notes(city, cfg)
    payload["units"] = {
        "temp": "°F" if (cfg["data"].get("temp_unit") or "C").upper() == "F" else "°C",
        "rainfall": city.rain_unit_label(),
        "raindays": "天",
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def write_xlsx(path: Path, city: CityClimate, cfg: dict[str, Any]) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("写出 XLSX 需要 openpyxl，请先安装：pip install openpyxl") from exc

    xcfg = cfg["table"].get("xlsx") or {}
    header, labels, cells = build_matrix(city, cfg)
    notes = build_notes(city, cfg)

    wb = Workbook()
    ws = wb.active
    ws.title = str(xcfg.get("sheet_name", "气候数据"))[:31]

    thin = Side(style="thin", color=str(xcfg.get("border_color", "FFBCC9D4")))
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=str(xcfg.get("header_fill", "FF0EBFA2")))
    zebra_fill = PatternFill("solid", fgColor=str(xcfg.get("zebra_fill", "FFF3F7FA")))
    header_font = Font(bold=True, color=str(xcfg.get("header_font_color", "FFFFFFFF")))
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    r = 1
    if cfg["table"].get("show_title_row", True):
        ws.cell(row=r, column=1, value=table_title(city, cfg)).font = Font(bold=True, size=13)
        r += 1
        for note in notes:
            ws.cell(row=r, column=1, value=note).font = Font(size=9, color="FF7A8899")
            r += 1
        r += 1

    header_row = r
    for c, value in enumerate(header, start=1):
        cell = ws.cell(row=header_row, column=c, value=value)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    num_fmt = str(xcfg.get("number_format", "0.0"))
    for i, (label, values) in enumerate(zip(labels, cells), start=1):
        row_idx = header_row + i
        cell = ws.cell(row=row_idx, column=1, value=label)
        cell.alignment = left
        cell.border = border
        if i % 2 == 0:
            cell.fill = zebra_fill
        for c, text in enumerate(values, start=2):
            value_cell = ws.cell(row=row_idx, column=c, value=text)
            value_cell.alignment = center
            value_cell.border = border
            if i % 2 == 0:
                value_cell.fill = zebra_fill
            if text and text != cfg["table"].get("missing_placeholder", "—"):
                value_cell.number_format = num_fmt

    width = float(xcfg.get("column_width", 13))
    ws.column_dimensions["A"].width = float(xcfg.get("first_column_width", 20))
    for c in range(2, len(header) + 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    if xcfg.get("freeze_header", True):
        ws.freeze_panes = ws.cell(row=header_row + 1, column=2)

    wb.save(path)
    return path


WRITERS = {
    "csv": (write_csv, ".csv"),
    "md": (write_markdown, ".md"),
    "markdown": (write_markdown, ".md"),
    "json": (write_json, ".json"),
    "xlsx": (write_xlsx, ".xlsx"),
}


def write_tables(city: CityClimate, cfg: dict[str, Any], out_dir: Path,
                 basename: str, logger=None) -> list[Path]:
    """按配置写出全部表格格式，返回生成的文件路径列表。"""
    formats = [f.lower() for f in (cfg["output"].get("table_formats") or [])]
    written: list[Path] = []
    for fmt in formats:
        writer_entry = WRITERS.get(fmt)
        if writer_entry is None:
            if logger:
                logger.warning(f"跳过不支持的表格格式：{fmt}")
            continue
        writer, ext = writer_entry
        path = out_dir / f"{basename}{ext}"
        writer(path, city, cfg)
        written.append(path)
        if logger:
            logger.debug(f"已写出表格：{path}")
    return written
